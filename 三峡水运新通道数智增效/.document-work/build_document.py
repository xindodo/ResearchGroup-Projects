from pathlib import Path
import re, json, shutil
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START, WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parent.parent
WORK=ROOT/'.document-work'
README=(ROOT/'README.md').read_text(encoding='utf-8-sig')
PROPOSAL=(WORK/'申报表述.md').read_text(encoding='utf-8-sig')
SOURCE=(ROOT/'参考资料/子课题三_表1表2检索依据.md').read_text(encoding='utf-8-sig')
ROWS=json.loads((WORK/'reference_rows.json').read_text(encoding='utf-8-sig'))
IMAGE=ROOT/'参考资料/子课题3_技术路线图_最新.png'
shutil.copyfile(ROOT/'.presentation-work/final-2.png',IMAGE)

def part(begin,end): return README.split(begin,1)[1].split(end,1)[0].strip()
context=part('## 一、项目及研究任务','## 五、本人主要研究任务')
context=context.replace('## 二、项目1.5主要研究内容及工作基础','### 1.2 项目1.5主要研究内容及工作基础').replace('## 三、课题一研究内容','### 1.3 课题一研究内容').replace('## 四、课题一子课题设置','### 1.4 课题一子课题设置')
context='### 1.1 项目与本人承担任务\n\n'+context
context=re.sub(r'^### ([1-5])\. (.+)$',r'#### 1.2.\1 \2',context,flags=re.M)
details='### 5.2'+part('### 5.2','### 5.9')
details=re.sub(r'^### 5\.([2-8])',lambda m:'### 4.'+str(int(m.group(1))-1),details,flags=re.M)
boundary='### 5.9'+part('### 5.9','## 六、')
boundary=boundary.replace('### 5.9','### 5.1')
relations=part('## 六、子课题三与另外四个子课题的联系和区别','## 七、')
relations=re.sub(r'^### 6\.([1-6])',lambda m:'### 5.'+str(int(m.group(1))+1),relations,flags=re.M)
concise=README.split('## 七、子课题三建议采用的申报表述',1)[1].strip()
ref=SOURCE.split('## 表1国外机构',1)[1]
ref='### 7.1 国外机构\n\n'+ref.replace('## 表2国内机构','### 7.2 国内机构')
ref=re.sub(r'^### (\d+)\. (.+)$',r'#### 原表序号\1 \2',ref,flags=re.M)
for row in ROWS[0]['rows']:
    # Preserve the comparative assessment added to original Table 1.
    needle='#### 原表序号'+row[0]+' '+row[1]
    start=ref.find(needle)
    if start>=0:
        nxt=ref.find('\n#### ',start+len(needle))
        next_section=ref.find('\n### 7.2',start+len(needle))
        if next_section>=0 and (nxt<0 or next_section<nxt):nxt=next_section
        if nxt<0:nxt=len(ref)
        ref=ref[:nxt]+'\n本项目研究内容自评价：'+row[5]+'\n'+ref[nxt:]

md='''# 子课题三研究方案与申报材料汇编

本子课题隶属于“三峡水运新通道数智增效”项目1.5的课题一，研究枢纽施工扰动对航运运行、运输服务及直接市场经营的动态影响，形成供“枢纽施工—航运波动—产业链响应”级联模型调用的航运响应模块。本文汇集项目背景、任务边界、申报表述、详细研究设计、子课题衔接、技术路线与国内外研究参考，供研究分工和申报材料编制使用。

## 一 项目及课题设置

'''+context+'\n\n'+PROPOSAL+'\n\n## 四 子课题三详细研究设计\n\n'+details+'\n\n## 五 研究边界与子课题衔接\n\n'+boundary+'\n\n'+relations+'\n\n### 5.8 综合申报表述\n\n'+concise+'''

## 六 子课题三技术路线

技术路线采用“研究基础与数据支撑—三项研究流程—成果输出与集成支撑”的总体结构。施工事件、通航约束及多源运行市场数据共同驱动响应识别；运行变化引起成本和运价调整，主体行为再反馈至有效运力、到达流与运输需求；以网络传播、恢复规律和参数校准为基础构建航运动态响应模块，支撑子课题五的级联模型集成与验证。

![子课题三技术路线图](参考资料/子课题3_技术路线图_最新.png)

可编辑图件：[子课题3 技术路线图](参考资料/子课题3_技术路线图.pptx)。文档第2页为最新方案，第1页保留原图。

## 七 国内外研究参考与表1表2续填内容

以下收录与子课题三相关的国外5家、国内5家机构研究资料及来源，对应申报书模板表1原序号16至20、表2原序号20至24。记录分别标明案例校准、历史数据试验、统计监测或工程应用，外部机构成果不作为本团队成果。正式申报时，按模板要求对每张整表遴选不超过5家机构。资料检索日期为2026年10月5日。

'''+ref+'''

## 八 项目资料索引

| 资料 | 用途 |
| --- | --- |
| README.md | 项目背景 课题设置与详细研究说明 |
| 参考资料/项目1.5申报书模板.docx | 表1与表2已续填 新增内容为红字 |
| 参考资料/子课题三_表1表2检索依据.md | 机构研究与应用证据 来源及检索记录 |
| 参考资料/子课题3_技术路线图.pptx | 原图与新增可编辑技术路线图 |
| 参考资料/子课题3_技术路线图_追加前备份.pptx | 新增路线图之前的原文档备份 |
| 子课题三_研究方案与申报材料汇编.md | 本汇编的Markdown版本 |

实施方案需进一步确定重点货类、实际施工情景、数据获取安排、时间与空间尺度、模块接口及考核指标。指标数值宜依据数据条件与先导试验确定，避免将外部研究的案例效果直接作为本项目承诺。
'''
# Remove numbering punctuation from heading labels while preserving official topic names.
md=re.sub(r'^(#{1,6}) (.+)$',lambda m:m[1]+' '+m[2].replace('：',' ').replace('、',' ').replace('——',' ').replace('“','').replace('”',''),md,flags=re.M)
OUTMD=ROOT/'子课题三_研究方案与申报材料汇编.md'
OUTDOC=ROOT/'子课题三_研究方案与申报材料汇编.docx'
OUTMD.write_text(md,encoding='utf-8')

doc=Document(); sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(2.1);sec.bottom_margin=Cm(2.1);sec.left_margin=Cm(2.2);sec.right_margin=Cm(2.2)
for name,size in [('Normal',10.5),('Title',20),('Heading 1',15),('Heading 2',12.5),('Heading 3',11.5),('Heading 4',11)]:
    sty=doc.styles[name];sty.font.name='宋体';sty.font.size=Pt(size);sty.font.color.rgb=RGBColor(0,0,0)
    sty.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'宋体')
    pf=sty.paragraph_format;pf.space_after=Pt(5);pf.line_spacing=1.25
    if name.startswith('Heading'):
        sty.font.bold=True;pf.space_before=Pt(10);pf.keep_with_next=True
        sty.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'黑体')
doc.styles['Normal'].paragraph_format.widow_control=True
doc.styles['Title'].paragraph_format.space_after=Pt(14)

def hyperlink(p,label,url):
    link=OxmlElement('w:hyperlink')
    rid=p.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    link.set(qn('r:id'),rid); r=OxmlElement('w:r'); rp=OxmlElement('w:rPr')
    fonts=OxmlElement('w:rFonts');fonts.set(qn('w:eastAsia'),'宋体');rp.append(fonts)
    color=OxmlElement('w:color');color.set(qn('w:val'),'0563C1');rp.append(color)
    r.append(rp);t=OxmlElement('w:t');t.text=label;r.append(t);link.append(r);p._p.append(link)
def inline(p,text):
    pat=r'(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^\)]+\))'
    for item in re.split(pat,text):
        if not item:continue
        if item.startswith('**'):
            p.add_run(item[2:-2]).bold=True
        elif item.startswith('`'):p.add_run(item[1:-1])
        elif item.startswith('['):
            m=re.match(r'\[([^\]]+)\]\(([^\)]+)\)',item)
            hyperlink(p,m[1],m[2] if m[2].startswith('http') else (ROOT/m[2]).resolve().as_uri())
        else:p.add_run(item)
def table(lines):
    vals=[[c.strip() for c in x.strip().strip('|').split('|')] for x in lines]
    vals=[r for r in vals if not all(re.fullmatch(r'[:\- ]+',c) for c in r)]
    t=doc.add_table(rows=len(vals),cols=len(vals[0]));t.autofit=False
    n=len(vals[0]); widths={2:[5,11.6],3:[2.0,10.3,4.3],4:[3.0,4.0,4.0,5.6]}.get(n,[16.6/n]*n)
    for c,w in zip(t.columns,widths):c.width=Cm(w)
    for i,(r,v) in enumerate(zip(t.rows,vals)):
        pr=r._tr.get_or_add_trPr()
        no=OxmlElement('w:cantSplit');pr.append(no)
        if i==0:pr.append(OxmlElement('w:tblHeader'))
        for j,(c,text) in enumerate(zip(r.cells,v)):
            c.width=Cm(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=c._tc.get_or_add_tcPr(); borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            cp.append(borders);marg=OxmlElement('w:tcMar')
            for side,val in [('top',80),('left',100),('bottom',80),('right',100)]:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),str(val));e.set(qn('w:type'),'dxa');marg.append(e)
            cp.append(marg);fill=OxmlElement('w:shd');fill.set(qn('w:val'),'clear');fill.set(qn('w:fill'),'E7EEF5' if i==0 else ('F8FAFC' if i%2==0 else 'FFFFFF'));cp.append(fill)
            order=['cnfStyle','tcW','gridSpan','hMerge','vMerge','tcBorders','shd','noWrap','tcMar','textDirection','tcFitText','vAlign','hideMark']
            children=sorted(list(cp),key=lambda x:order.index(x.tag.split('}')[-1]) if x.tag.split('}')[-1] in order else 999)
            for node in children:cp.remove(node)
            for node in children:cp.append(node)
            p=c.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.15
            inline(p,text)
            for run in p.runs:run.font.size=Pt(9);run.bold=run.bold or i==0
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

lines=md.splitlines();i=0
while i<len(lines):
    line=lines[i].strip()
    if not line:i+=1;continue
    if line.startswith('|'):
        group=[]
        while i<len(lines) and lines[i].strip().startswith('|'):group.append(lines[i]);i+=1
        table(group);continue
    if line.startswith('!['):
        landscape=doc.add_section(WD_SECTION_START.NEW_PAGE)
        landscape.orientation=WD_ORIENT.LANDSCAPE;landscape.page_width=Cm(29.7);landscape.page_height=Cm(21)
        landscape.top_margin=Cm(1.6);landscape.bottom_margin=Cm(1.6);landscape.left_margin=Cm(1.5);landscape.right_margin=Cm(1.5)
        hp=doc.add_paragraph('子课题三技术路线图',style='Heading 1');hp.paragraph_format.space_before=Pt(0)
        p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        run=p.add_run();run.add_picture(str(IMAGE),width=Cm(26.5))
        pic=run._r.xpath('.//wp:docPr')[0];pic.set('descr','枢纽施工扰动下航运系统波动响应研究技术路线图 展示数据支撑 三项研究内容 行为反馈 方法 成果及研究边界')
        cap=doc.add_paragraph('图1 子课题三最新技术路线图');cap.alignment=WD_ALIGN_PARAGRAPH.CENTER;cap.paragraph_format.space_after=Pt(0)
        portrait=doc.add_section(WD_SECTION_START.NEW_PAGE);portrait.orientation=WD_ORIENT.PORTRAIT
        portrait.page_width=Cm(21);portrait.page_height=Cm(29.7);portrait.top_margin=Cm(2.1);portrait.bottom_margin=Cm(2.1);portrait.left_margin=Cm(2.2);portrait.right_margin=Cm(2.2)
        i+=1;continue
    if line=='\\[':
        eq=[];i+=1
        while i<len(lines) and lines[i].strip()!='\\]':eq.append(lines[i]);i+=1
        p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        om=OxmlElement('m:oMath');mr=OxmlElement('m:r');mt=OxmlElement('m:t');mt.text='Qₜ₊₁ = max{0, Qₜ + Aₜ − Cₜ}';mr.append(mt);om.append(mr);p._p.append(om)
        i+=1;continue
    m=re.match(r'^(#{1,6}) (.+)$',line)
    if m:
        level=len(m[1]);p=doc.add_paragraph(style='Title' if level==1 else 'Heading '+str(min(level-1,4)))
        inline(p,m[2])
    else:
        line=re.sub(r'^>\s*','',line)
        p=doc.add_paragraph();inline(p,line)
        if line.startswith('**') and line.endswith('**') and len(line)<60:p.paragraph_format.keep_with_next=True
        if len(line)>70:p.paragraph_format.first_line_indent=Pt(21)
    i+=1
doc.core_properties.title='子课题三研究方案与申报材料汇编'
doc.core_properties.subject='枢纽施工扰动下航运系统波动响应研究'
doc.core_properties.author='';doc.core_properties.keywords='三峡 枢纽施工 航运波动 市场反馈 级联影响'
doc.save(OUTDOC)
print(json.dumps({'docx':str(OUTDOC),'markdown':str(OUTMD),'chars':len(md),'paragraphs':len(doc.paragraphs),'tables':len(doc.tables),'images':len(doc.inline_shapes)},ensure_ascii=False))
